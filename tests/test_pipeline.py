from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.pipeline import process_frame, read_main_export, write_result


def sample_frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "DateConverted Hierarchy - Date": [
                "2026-01-01",
                "2026-01-01",
                "2026-01-02",
                "bad-date",
            ],
            "Time": ["08:00", "08:05", "08:00", "08:10"],
            "Green Time (Sec)": [100, 49, 205, "bad"],
            "GT_mean": [100, 100, 100, 100],
        }
    )


def test_process_frame_builds_report_and_daily_frames():
    result = process_frame(sample_frame(), source_name="sanitized.xlsx")
    report = result.report.iloc[0]

    assert report["Rows"] == 4
    assert report["GT_Valid"] == 3
    assert report["GT_Invalid"] == 1
    assert report["GT_Anomalies"] == 2
    assert report["GT_Normal"] == 1
    assert report["GT_Anomaly_Rate"] == pytest.approx(2 / 3)
    assert list(result.daily_frames) == ["2026-01-01", "2026-01-02"]
    assert len(result.daily_frames["2026-01-01"]) == 2


def test_process_frame_without_date_column_still_processes():
    frame = sample_frame().drop(columns=["DateConverted Hierarchy - Date"])
    result = process_frame(frame)
    assert result.daily_frames == {}
    assert len(result.processed) == 4


def test_empty_export_with_required_schema_is_safe():
    frame = pd.DataFrame(columns=["Green Time (Sec)", "GT_mean"])
    result = process_frame(frame, source_name="empty.csv")
    report = result.report.iloc[0]

    assert len(result.processed) == 0
    assert report["Rows"] == 0
    assert report["GT_Valid"] == 0
    assert report["GT_Invalid"] == 0
    assert report["GT_Anomalies"] == 0
    assert report["GT_Anomaly_Rate"] == 0.0
    assert result.daily_frames == {}


def test_write_result_creates_auditable_workbook(tmp_path: Path):
    result = process_frame(sample_frame(), source_name="sanitized.xlsx")
    output = tmp_path / "screened.xlsx"
    write_result(result, output)

    assert output.exists()
    workbook = pd.ExcelFile(output, engine="openpyxl")
    assert workbook.sheet_names == [
        "RUN_REPORT",
        "PROCESSED",
        "2026-01-01",
        "2026-01-02",
    ]

    report = pd.read_excel(output, sheet_name="RUN_REPORT", engine="openpyxl")
    processed = pd.read_excel(output, sheet_name="PROCESSED", engine="openpyxl")
    assert report.loc[0, "GT_Anomalies"] == 2
    assert "GT_Reason" in processed.columns

    styled = load_workbook(output)
    assert all(sheet.freeze_panes == "A2" for sheet in styled.worksheets)
    assert all(sheet.auto_filter.ref for sheet in styled.worksheets)
    rate_column = next(
        cell.column
        for cell in styled["RUN_REPORT"][1]
        if cell.value == "GT_Anomaly_Rate"
    )
    assert styled["RUN_REPORT"].cell(2, rate_column).number_format == "0.0%"


def test_read_main_export_supports_csv_and_xlsx(tmp_path: Path):
    frame = sample_frame()
    csv_path = tmp_path / "main.csv"
    xlsx_path = tmp_path / "main.xlsx"
    frame.to_csv(csv_path, index=False)
    frame.to_excel(xlsx_path, index=False, engine="openpyxl")

    assert len(read_main_export(csv_path)) == 4
    assert len(read_main_export(xlsx_path)) == 4


def test_read_main_export_rejects_unknown_format(tmp_path: Path):
    path = tmp_path / "main.txt"
    path.write_text("not supported", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported MAIN input type"):
        read_main_export(path)
