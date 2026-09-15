from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

from .config import MAIN_DATE_COLUMN, GTThresholds, DEFAULT_THRESHOLDS
from .data_quality import summarize_main_data_quality
from .gt_anomaly import add_gt_anomaly_columns


@dataclass(frozen=True)
class ProcessingResult:
    processed: pd.DataFrame
    report: pd.DataFrame
    daily_frames: dict[str, pd.DataFrame]


def _safe_sheet_name(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", value).strip()
    return (cleaned or fallback)[:31]


def _daily_frames(processed: pd.DataFrame, date_column: str) -> dict[str, pd.DataFrame]:
    if date_column not in processed.columns:
        return {}

    parsed_dates = pd.to_datetime(processed[date_column], errors="coerce")
    frames: dict[str, pd.DataFrame] = {}
    for date in sorted(parsed_dates.dropna().dt.date.unique()):
        mask = parsed_dates.dt.date == date
        key = _safe_sheet_name(str(date), "DAY")
        frames[key] = processed.loc[mask].copy()
    return frames


def build_run_report(processed: pd.DataFrame, source_name: str) -> pd.DataFrame:
    anomaly = processed["GT_Anomaly"]
    severity = processed["GT_Severity"]
    valid = anomaly != "Invalid"

    row = {
        "Source_MAIN": source_name,
        "Rows": int(len(processed)),
        "GT_Valid": int(valid.sum()),
        "GT_Invalid": int((anomaly == "Invalid").sum()),
        "GT_Anomalies": int((anomaly == "Anomaly").sum()),
        "GT_Normal": int((anomaly == "No Anomaly").sum()),
        "Severe_Low": int((severity == "Severe-Low").sum()),
        "Anomaly_Low": int((severity == "Anomaly-Low").sum()),
        "Anomaly_High": int((severity == "Anomaly-High").sum()),
        "Severe_High": int((severity == "Severe-High").sum()),
        **summarize_main_data_quality(processed),
    }
    row["GT_Anomaly_Rate"] = (
        row["GT_Anomalies"] / row["GT_Valid"] if row["GT_Valid"] else 0.0
    )
    return pd.DataFrame([row])


def process_frame(
    frame: pd.DataFrame,
    *,
    source_name: str = "in_memory",
    date_column: str = MAIN_DATE_COLUMN,
    thresholds: GTThresholds = DEFAULT_THRESHOLDS,
) -> ProcessingResult:
    processed = add_gt_anomaly_columns(frame, thresholds)
    report = build_run_report(processed, source_name)
    return ProcessingResult(
        processed=processed,
        report=report,
        daily_frames=_daily_frames(processed, date_column),
    )


def read_main_export(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return pd.read_excel(path, engine="openpyxl")
    if suffix == ".csv":
        return pd.read_csv(path)
    raise ValueError(f"Unsupported MAIN input type: {suffix or '<none>'}; use .xlsx or .csv")


def _format_workbook(output_path: Path) -> None:
    workbook = load_workbook(output_path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions

        for column_index in range(1, sheet.max_column + 1):
            values = [sheet.cell(row=row, column=column_index).value for row in range(1, sheet.max_row + 1)]
            width = max((len(str(value)) for value in values if value is not None), default=0)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 12), 40)

    report = workbook["RUN_REPORT"]
    headers = {cell.value: cell.column for cell in report[1]}
    anomaly_rate_column = headers.get("GT_Anomaly_Rate")
    if anomaly_rate_column is not None and report.max_row >= 2:
        report.cell(row=2, column=anomaly_rate_column).number_format = "0.0%"

    workbook.save(output_path)


def write_result(result: ProcessingResult, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        result.report.to_excel(writer, index=False, sheet_name="RUN_REPORT")
        result.processed.to_excel(writer, index=False, sheet_name="PROCESSED")
        for sheet_name, frame in result.daily_frames.items():
            candidate = sheet_name
            if candidate in {"RUN_REPORT", "PROCESSED"}:
                candidate = f"DAY_{candidate}"[:31]
            frame.to_excel(writer, index=False, sheet_name=candidate)

    _format_workbook(output_path)
